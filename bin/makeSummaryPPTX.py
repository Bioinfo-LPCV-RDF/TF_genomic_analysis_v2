







import argparse #v1.1
from pptx import Presentation #v0.6.21
from pptx.util import Inches, Pt, Cm
import os

parser = argparse.ArgumentParser()

parser.add_argument("--pwm", "-p",nargs='+',type = str)
parser.add_argument("--pwmdir", "-pd",nargs='+',type = str)
parser.add_argument("--meme", "-m",nargs='+',type = str)
parser.add_argument("--kmer", "-k",nargs='+',type = str)
parser.add_argument("--tffm", "-t",nargs='+',type = str)
parser.add_argument("--roc", "-r",nargs='+',type = str)
parser.add_argument("--rocsall", "-rall",nargs='+',type = str)
parser.add_argument("--spacing", "-s",nargs='+',type = str)
parser.add_argument("--deciles", "-dc",nargs='+',type = str)
parser.add_argument("--numberpeaks", "-pknb", nargs='+', help='')
parser.add_argument("--comparisons", "-c",nargs='+',type = str)
parser.add_argument("--names", "-n",nargs='+',type = str)
parser.add_argument("--namescomps", "-nc",nargs='+',type = str)
parser.add_argument("--replicatesComparisons", "-rc",nargs='+',type = str)
parser.add_argument("--out", "-o", help='path and filename with extension')




args = parser.parse_args()

def main(ROCs, ALLROCS, PWMs, dirPWM, MEME, Kmer, Tffm, Spacings, Names, Namescomps, ReplicatesComparisons, comparisons, deciles, out, numberpeaks):
	# print(numberpeaks)
	# print(dirPWM)
	startingMEME=0
	# Create a new PowerPoint presentation
	prs = Presentation()
	for ROC,motifPWM,spacing,replicateComparison,samplename,peakNumber,PWMdirectory,AllRoc, Onekmer, Onetffm  in zip(ROCs,PWMs,Spacings,ReplicatesComparisons,Names,numberpeaks,dirPWM,ALLROCS, Kmer, Tffm):
		slide = prs.slides.add_slide(prs.slide_layouts[6])
		textbox = slide.shapes.add_textbox(Cm(10), 0, 200, 100)
		textbox.text = "Basic Analysis:"+samplename+"\n Number of Consensus peaks: "+str(peakNumber)
		i=1
		while os.path.exists(PWMdirectory+"/logo"+str(i)+".png"):
			i+=1
		i-=1
		savei=i
		while i>0:
			
			PWM1=slide.shapes.add_picture(PWMdirectory+"/logo"+str(i)+".png", left=Cm(4*(i-1)), top=Cm(13.4), height=Cm(2.5))
			PWM2=slide.shapes.add_picture(PWMdirectory+"/logo_rc"+str(i)+".png", left=Cm(4*(i-1)), top=Cm(16), height=Cm(2.5))
			textbox = slide.shapes.add_textbox(Cm((4*(i-1))+1), Cm(12.8), 200, 100)
			if int(peakNumber)>=600:
				textbox.text = str(MEME[startingMEME+i-1])+"/600"
			else:
				textbox.text = str(MEME[startingMEME+i-1])+"/"+str(peakNumber)
			i-=1
		startingMEME+=savei
		ROCall=slide.shapes.add_picture(AllRoc, left=Cm(19.6), top=Cm(13.5), height= Cm(5.6))
		if replicateComparison != "NA":
			picREP = slide.shapes.add_picture(replicateComparison, left=0, top=Cm(4), height=Cm(8.44))
		else:
			textbox = slide.shapes.add_textbox(Cm(2), Cm(4), 200, 100)
			textbox.text = "No replicates for this sample"

		# PWM1=slide.shapes.add_picture(PWMdirectory+"/logo1.png", left=0, top=Cm(3), height=Cm(3))
		# PWM2=slide.shapes.add_picture(PWMdirectory+"/logo2.png", left=0, top=Cm(6), height=Cm(3))
		# PWM3=slide.shapes.add_picture(PWMdirectory+"/logo3.png", left=0, top=Cm(9), height=Cm(3))
		# PWM4=slide.shapes.add_picture(PWMdirectory+"/logo4.png", left=0, top=Cm(12), height=Cm(3))
		# PWM5=slide.shapes.add_picture(PWMdirectory+"/logo5.png", left=0, top=Cm(15), height=Cm(3))



		# Add a slide to the presentation
		slide = prs.slides.add_slide(prs.slide_layouts[6])
		topPWM = Cm(8); heightPWM = Cm(4)
		topROC = Cm(12); heightROC = Cm(7)
		leftSPA = Cm(12); topSPA = Cm(7); heightSPA = Cm(12)
	
		# tffm1=slide.shapes.add_picture(Onetffm, left=Cm(0), top=Cm(4), height=Cm(2.5))
		kmer1=slide.shapes.add_picture(Onekmer, left=Cm(0), top=Cm(5.5), height=Cm(2.5))

		picPWM = slide.shapes.add_picture(motifPWM, left=0, top=topPWM, height=heightPWM)
		picROC = slide.shapes.add_picture(ROC, left=0, top=topROC,height=heightROC)
		picSPA = slide.shapes.add_picture(spacing, left=leftSPA, top=topSPA, height=heightSPA)

		## Add a textbox to the slide
		textbox = slide.shapes.add_textbox(Cm(10), 0, 200, 100)

		## Add text to the textbox
		textbox.text = "Basic Analysis:"+samplename+"\n Number of Consensus peaks: "+str(peakNumber)
	for comp,decile,namecomp in zip(comparisons,deciles,Namescomps):
		slide = prs.slides.add_slide(prs.slide_layouts[6])
		
		textbox = slide.shapes.add_textbox(Cm(10), 0, 200, 100)
		textbox.text = "comparison: "+" ".join(namecomp.split("_"))
		
		picComp = slide.shapes.add_picture(comp, left=0, top=Cm(2), height=Cm(10))
		picdecile = slide.shapes.add_picture(decile, left=Cm(15.3), top=Cm(7.69), height=Cm(11.36))

	# Save the PowerPoint presentation
	prs.save(out)

main(args.roc, args.rocsall, args.pwm, args.pwmdir, args.meme, args.kmer, args.tffm, args.spacing, args.names, args.namescomps, args.replicatesComparisons, args.comparisons, args.deciles, args.out, args.numberpeaks)


 